"""레거시 .xls 업로드 지원 — .xls로 저장한 픽스처가 .xlsx와 똑같이 파싱되는지."""
import datetime
from io import BytesIO
from pathlib import Path

import openpyxl
import pytest
import xlwt

from app.parsers import ParseError, detect
from app.parsers.xls import is_xls, xls_to_xlsx

FIXTURES = Path(__file__).resolve().parents[1] / "fixtures"


def xlsx_to_xls(xlsx_bytes: bytes) -> bytes:
    """테스트용: .xlsx 픽스처를 실제 BIFF .xls로 다시 저장한다."""
    src = openpyxl.load_workbook(BytesIO(xlsx_bytes), data_only=True)
    book = xlwt.Workbook()
    date_style = xlwt.easyxf(num_format_str="YYYY-MM-DD HH:MM:SS")
    for ws in src.worksheets:
        sheet = book.add_sheet(ws.title)
        for r, row in enumerate(ws.iter_rows(values_only=True)):
            for c, v in enumerate(row):
                if v is None:
                    continue
                if isinstance(v, (datetime.datetime, datetime.date)):
                    sheet.write(r, c, v, date_style)
                elif isinstance(v, datetime.time):
                    sheet.write(r, c, v.strftime("%H:%M:%S"))
                else:
                    sheet.write(r, c, v)
    out = BytesIO()
    book.save(out)
    return out.getvalue()


@pytest.mark.parametrize(
    "fixture", ["samsung-card-fixture", "woori-card-fixture", "hana-bank-fixture"]
)
def test_xls_parses_same_as_xlsx(fixture):
    xlsx = (FIXTURES / f"{fixture}.xlsx").read_bytes()
    xls = xlsx_to_xls(xlsx)
    assert is_xls(xls) and not is_xls(xlsx)

    converted = xls_to_xlsx(xls)
    st_a, parse_a = detect(converted)
    st_b, parse_b = detect(xlsx)
    assert st_a == st_b

    a, b = parse_a(converted), parse_b(xlsx)
    assert len(a.transactions) == len(b.transactions) > 0
    for x, y in zip(a.transactions, b.transactions, strict=True):
        assert (x.txn_date, x.amount, x.merchant_raw) == (y.txn_date, y.amount, y.merchant_raw)


def test_xls_to_xlsx_rejects_garbage():
    with pytest.raises(ParseError) as exc:
        xls_to_xlsx(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 100)
    assert exc.value.code == "WORKBOOK_LOAD_FAILED"
