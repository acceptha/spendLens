"""레거시 .xls(BIFF/OLE2) → .xlsx 바이트 변환.

카드사·은행 사이트가 내려주는 명세서는 .xls인 경우가 많다. 파서는 전부 openpyxl(.xlsx)
기준이므로, 업로드 시 .xls면 메모리에서 .xlsx로 옮겨 담아 기존 파서를 그대로 쓴다.
날짜 셀은 datetime으로 변환해 .xlsx에서 읽은 것과 같은 타입이 되게 한다.
"""
from io import BytesIO

import openpyxl
import xlrd

from app.parsers import ParseError

OLE2_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def is_xls(file_bytes: bytes) -> bool:
    return file_bytes.startswith(OLE2_MAGIC)


def xls_to_xlsx(file_bytes: bytes) -> bytes:
    try:
        book = xlrd.open_workbook(file_contents=file_bytes)
    except Exception as e:  # noqa: BLE001 — xlrd는 다양한 예외를 던진다
        raise ParseError("WORKBOOK_LOAD_FAILED", reason=str(e)) from e

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet in book.sheets():
        ws = wb.create_sheet(title=sheet.name[:31] or "Sheet")
        for r in range(sheet.nrows):
            row = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    row.append(None)
                elif cell.ctype == xlrd.XL_CELL_NUMBER and float(cell.value).is_integer():
                    # xlrd는 숫자를 모두 float로 준다. .xlsx(openpyxl)처럼 정수는 int로.
                    row.append(int(cell.value))
                else:
                    row.append(cell.value)
            ws.append(row)

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
